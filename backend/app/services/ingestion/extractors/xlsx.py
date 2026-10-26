from __future__ import annotations

import math
import time
from datetime import date, datetime
from datetime import time as time_value
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.observability import record_document_extraction
from app.services.ingestion.errors import (
    ExtractionError,
    ExtractionErrorCode,
    map_unexpected_extraction_failure,
)
from app.services.ingestion.extractors._common import (
    TrustedExtractionPolicy,
    append_warning_once,
    clean_contract_text,
    preflight_office_package,
    sha256_file,
    validate_trusted_input_path,
)
from app.services.ingestion.types import (
    ConfirmedDocumentFormat,
    ExtractedBlock,
    ExtractedBlockType,
    ExtractionLimits,
    ExtractionMethod,
    ExtractionResult,
    ExtractionStatistics,
    ExtractionWarningCode,
)

if TYPE_CHECKING:
    from app.core.config import Settings

XLSX_EXTRACTOR_NAME = "openpyxl"
XLSX_EXTRACTOR_VERSION = "1.0.0"


class XlsxExtractor:
    format_name = ConfirmedDocumentFormat.XLSX
    name = XLSX_EXTRACTOR_NAME
    version = XLSX_EXTRACTOR_VERSION

    def __init__(self, policy: TrustedExtractionPolicy) -> None:
        self._policy = policy

    @classmethod
    def from_settings(cls, settings: Settings) -> XlsxExtractor:
        return cls(TrustedExtractionPolicy.from_settings(settings))

    def supports(self, confirmed_format: ConfirmedDocumentFormat) -> bool:
        return confirmed_format is ConfirmedDocumentFormat.XLSX

    def extract(self, path: Path, limits: ExtractionLimits) -> ExtractionResult:
        started = time.monotonic()
        workbook = None
        try:
            trusted_path = validate_trusted_input_path(
                path,
                processing_root=self._policy.processing_root,
                expected_suffix=".xlsx",
                maximum_bytes=self._policy.maximum_input_bytes,
            )
            inspection = preflight_office_package(
                trusted_path,
                confirmed_suffix=".xlsx",
                limits=limits,
            )
            source_sha256 = sha256_file(
                trusted_path,
                maximum_bytes=self._policy.maximum_input_bytes,
            )
            workbook = load_workbook(
                filename=trusted_path,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            result = _extract_workbook(
                workbook,
                limits=limits,
                source_sha256=source_sha256,
                dimensions_need_correction=inspection.worksheet_dimensions_need_correction,
                formula_without_cache=inspection.has_formula_without_cached_value,
            )
            _close_workbook(workbook)
            workbook = None
            record_document_extraction(
                document_format="xlsx",
                duration_seconds=time.monotonic() - started,
                outcome="success",
                block_count=result.statistics.block_count,
                row_count=result.statistics.row_count,
                worksheet_count=result.statistics.worksheet_count,
            )
            return result
        except ExtractionError as exc:
            record_document_extraction(
                document_format="xlsx",
                duration_seconds=time.monotonic() - started,
                outcome=_outcome(exc.code),
            )
            raise
        except (InvalidFileException, KeyError, TypeError, ValueError):
            record_document_extraction(
                document_format="xlsx",
                duration_seconds=time.monotonic() - started,
                outcome="rejected",
            )
            raise ExtractionError(ExtractionErrorCode.DOCUMENT_MALFORMED) from None
        except MemoryError:
            record_document_extraction(
                document_format="xlsx",
                duration_seconds=time.monotonic() - started,
                outcome="failed",
            )
            raise ExtractionError(ExtractionErrorCode.DOCUMENT_PARSER_RESOURCE_EXHAUSTED) from None
        except Exception as exc:
            record_document_extraction(
                document_format="xlsx",
                duration_seconds=time.monotonic() - started,
                outcome="failed",
            )
            raise map_unexpected_extraction_failure(exc) from None
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass


def _extract_workbook(
    workbook: object,
    *,
    limits: ExtractionLimits,
    source_sha256: str,
    dimensions_need_correction: bool,
    formula_without_cache: bool,
) -> ExtractionResult:
    blocks: list[ExtractedBlock] = []
    warnings: list[ExtractionWarningCode] = []
    worksheet_count = 0
    row_count = 0
    cell_count = 0
    character_count = 0

    if formula_without_cache:
        append_warning_once(
            warnings,
            ExtractionWarningCode.FORMULA_WITHOUT_CACHED_VALUE,
            limits.maximum_warnings,
        )

    for worksheet in workbook.worksheets:
        if worksheet.sheet_state != "visible":
            append_warning_once(
                warnings,
                ExtractionWarningCode.HIDDEN_SHEET_SKIPPED,
                limits.maximum_warnings,
            )
            continue
        worksheet_count += 1
        if worksheet_count > limits.maximum_worksheets:
            raise _limit_error()
        if dimensions_need_correction:
            worksheet.reset_dimensions()
            append_warning_once(
                warnings,
                ExtractionWarningCode.WORKSHEET_DIMENSIONS_CORRECTED,
                limits.maximum_warnings,
            )

        in_region = False
        for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_number > limits.maximum_rows_per_sheet:
                if any(value is not None for value in values):
                    raise _limit_error()
                break
            trimmed = list(values)
            while trimmed and trimmed[-1] is None:
                trimmed.pop()
            if not trimmed or all(value is None or value == "" for value in trimmed):
                in_region = False
                continue
            if len(trimmed) > limits.maximum_spreadsheet_columns:
                raise _limit_error()
            cell_count += sum(value is not None for value in trimmed)
            if cell_count > limits.maximum_spreadsheet_cells:
                raise _limit_error()
            rendered = tuple(_canonical_cell(value) for value in trimmed)
            text = " | ".join(rendered)
            if not text.strip(" |"):
                in_region = False
                continue
            is_header = not in_region
            in_region = True
            character_count = _append_row_block(
                blocks,
                text=text,
                sheet_name=worksheet.title,
                row_number=row_number,
                is_header=is_header,
                character_count=character_count,
                limits=limits,
            )
            row_count += 1

    if not blocks:
        raise ExtractionError(ExtractionErrorCode.DOCUMENT_NO_EXTRACTABLE_CONTENT)
    try:
        statistics = ExtractionStatistics(
            detected_format=ConfirmedDocumentFormat.XLSX,
            page_count=0,
            worksheet_count=worksheet_count,
            table_count=0,
            row_count=row_count,
            block_count=len(blocks),
            character_count=character_count,
            native_page_count=0,
            ocr_page_count=0,
            warning_codes=warnings,
            extractor_name=XLSX_EXTRACTOR_NAME,
            extractor_version=XLSX_EXTRACTOR_VERSION,
            limits=limits,
        )
        return ExtractionResult(
            blocks=blocks,
            statistics=statistics,
            source_content_sha256=source_sha256,
            confirmed_format=ConfirmedDocumentFormat.XLSX,
            limits=limits,
        )
    except ValueError:
        raise _limit_error() from None


def _canonical_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime | date | time_value):
        return value.isoformat()
    if isinstance(value, int | Decimal):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ExtractionError(ExtractionErrorCode.DOCUMENT_MALFORMED)
        return str(value)
    if isinstance(value, str):
        return clean_contract_text(value)
    raise ExtractionError(ExtractionErrorCode.DOCUMENT_MALFORMED)


def _append_row_block(
    blocks: list[ExtractedBlock],
    *,
    text: str,
    sheet_name: str,
    row_number: int,
    is_header: bool,
    character_count: int,
    limits: ExtractionLimits,
) -> int:
    if len(blocks) >= limits.maximum_blocks:
        raise _limit_error()
    new_character_count = character_count + len(text)
    if new_character_count > limits.maximum_extracted_characters:
        raise _limit_error()
    try:
        blocks.append(
            ExtractedBlock(
                ordinal=len(blocks),
                text=text,
                block_type=ExtractedBlockType.WORKSHEET_ROW,
                sheet_name=sheet_name,
                row_start=row_number,
                row_end=row_number,
                extraction_method=ExtractionMethod.NATIVE,
                metadata={"source_kind": "worksheet", "is_header": is_header},
                limits=limits,
            )
        )
    except ValueError:
        raise _limit_error() from None
    return new_character_count


def _limit_error() -> ExtractionError:
    return ExtractionError(ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED)


def _close_workbook(workbook: object) -> None:
    try:
        workbook.close()
    except Exception:
        raise ExtractionError(ExtractionErrorCode.DOCUMENT_EXTRACTION_FAILED) from None


def _outcome(code: ExtractionErrorCode) -> str:
    if code in {
        ExtractionErrorCode.DOCUMENT_ENCRYPTED,
        ExtractionErrorCode.DOCUMENT_MALFORMED,
        ExtractionErrorCode.UNSUPPORTED_DOCUMENT_FORMAT,
        ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED,
        ExtractionErrorCode.DOCUMENT_NO_EXTRACTABLE_CONTENT,
    }:
        return "rejected"
    return "failed"
