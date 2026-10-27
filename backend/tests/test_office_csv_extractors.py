from __future__ import annotations

import csv
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from extraction_fixtures import extraction_limits
from openpyxl import Workbook

from app.core.config import Settings
from app.services.ingestion.errors import ExtractionError, ExtractionErrorCode
from app.services.ingestion.extractors.csv import CSV_EXTRACTOR_NAME, CsvExtractor
from app.services.ingestion.extractors.docx import DOCX_EXTRACTOR_NAME, DocxExtractor
from app.services.ingestion.extractors.xlsx import XLSX_EXTRACTOR_NAME, XlsxExtractor
from app.services.ingestion.types import (
    ConfirmedDocumentFormat,
    ExtractedBlockType,
    ExtractionWarningCode,
)


def _trusted_path(tmp_path: Path, suffix: str) -> tuple[Path, Path]:
    root = tmp_path / "processing"
    workspace = root / str(uuid4())
    workspace.mkdir(parents=True)
    return root, workspace / f"input{suffix}"


def _settings(root: Path) -> Settings:
    return Settings(
        _env_file=None,
        document_processing_temp_root=str(root),
        max_upload_bytes=5_000_000,
        max_request_body_bytes=5_000_001,
    )


def test_docx_preserves_heading_paragraph_table_list_and_body_order(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    document.add_heading("Delivery", level=1)
    document.add_paragraph("Opening scope")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "Design"
    table.cell(1, 1).text = "Team A"
    document.add_paragraph("Approval required", style="List Bullet")
    document.add_paragraph("Closing scope")
    document.sections[0].header.paragraphs[0].text = "Excluded header"
    document.save(path)

    result = DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert [block.text for block in result.blocks] == [
        "Delivery",
        "Opening scope",
        "Item | Owner",
        "Design | Team A",
        "Approval required",
        "Closing scope",
    ]
    assert [block.block_type for block in result.blocks] == [
        ExtractedBlockType.HEADING,
        ExtractedBlockType.PARAGRAPH,
        ExtractedBlockType.TABLE_ROW,
        ExtractedBlockType.TABLE_ROW,
        ExtractedBlockType.LIST_ITEM,
        ExtractedBlockType.PARAGRAPH,
    ]
    assert result.blocks[-1].heading_path == ("Delivery",)
    assert result.statistics.table_count == 1
    assert result.statistics.row_count == 2
    assert result.statistics.extractor_name == DOCX_EXTRACTOR_NAME
    assert result.statistics.warning_codes == (ExtractionWarningCode.HEADER_FOOTER_EXCLUDED,)
    assert "Excluded header" not in str(result.to_jsonable())


def test_docx_merged_cells_are_not_duplicated(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Merged scope"
    document.save(path)

    result = DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert result.blocks[0].text == "Merged scope"


def test_docx_without_body_text_fails_closed(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    Document().save(path)

    with pytest.raises(ExtractionError) as error:
        DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_NO_EXTRACTABLE_CONTENT


@pytest.mark.parametrize(
    ("entry_name", "expected_code"),
    [
        ("../escape.xml", ExtractionErrorCode.DOCUMENT_MALFORMED),
        ("word/vbaProject.bin", ExtractionErrorCode.UNSUPPORTED_DOCUMENT_FORMAT),
        ("word/embeddings/object.bin", ExtractionErrorCode.UNSUPPORTED_DOCUMENT_FORMAT),
    ],
)
def test_docx_rejects_unsafe_archive_parts(
    tmp_path: Path,
    entry_name: str,
    expected_code: ExtractionErrorCode,
) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    document.add_paragraph("Synthetic scope")
    document.save(path)
    with ZipFile(path, "a", ZIP_DEFLATED) as package:
        package.writestr(entry_name, b"synthetic")

    with pytest.raises(ExtractionError) as error:
        DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is expected_code


def test_docx_rejects_xml_entities_without_exposing_content(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    document.add_paragraph("Synthetic scope")
    document.save(path)
    _replace_zip_part(
        path,
        "word/document.xml",
        b'<!DOCTYPE root [<!ENTITY unsafe "sensitive-marker">]><root>&unsafe;</root>',
    )

    with pytest.raises(ExtractionError) as error:
        DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_MALFORMED
    assert "sensitive-marker" not in str(error.value)


def test_docx_external_relationship_is_never_fetched(tmp_path: Path, monkeypatch) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    document.add_paragraph("Synthetic scope")
    document.part.relate_to(
        "https://external.invalid/resource",
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    document.save(path)

    def reject_network(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("network access attempted")

    monkeypatch.setattr("socket.create_connection", reject_network)
    result = DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert result.blocks[0].text == "Synthetic scope"


def test_docx_enforces_archive_entry_limit(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    document.add_paragraph("Synthetic scope")
    document.save(path)

    with pytest.raises(ExtractionError) as error:
        DocxExtractor.from_settings(_settings(root)).extract(
            path,
            extraction_limits(maximum_archive_entries=1),
        )

    assert error.value.code is ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED


def test_office_compound_encryption_container_is_rejected(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".docx")
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"synthetic")

    with pytest.raises(ExtractionError) as error:
        DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_ENCRYPTED


def test_xlsx_preserves_visible_sheet_rows_and_never_evaluates_formulas(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Scope"
    worksheet.append(["Item", "Amount"])
    worksheet.append(["Design", 12.5])
    worksheet.append([])
    worksheet.append(["Risk", "Owner"])
    worksheet.append(["Delay", "Team A"])
    worksheet["C2"] = "=1+1"
    hidden = workbook.create_sheet("Private")
    hidden.sheet_state = "hidden"
    hidden.append(["Must not appear"])
    workbook.save(path)
    workbook.close()

    result = XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert result.confirmed_format is ConfirmedDocumentFormat.XLSX
    assert [block.row_start for block in result.blocks] == [1, 2, 4, 5]
    assert result.blocks[0].metadata["is_header"] is True
    assert result.blocks[2].metadata["is_header"] is True
    assert all(block.sheet_name == "Scope" for block in result.blocks)
    assert "=1+1" not in str(result.to_jsonable())
    assert result.statistics.worksheet_count == 1
    assert result.statistics.extractor_name == XLSX_EXTRACTOR_NAME
    assert set(result.statistics.warning_codes) == {
        ExtractionWarningCode.FORMULA_WITHOUT_CACHED_VALUE,
        ExtractionWarningCode.HIDDEN_SHEET_SKIPPED,
    }


def test_xlsx_uses_cached_formula_value_without_evaluation(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Item", "Value"])
    worksheet.append(["Cached", "=1+1"])
    workbook.save(path)
    workbook.close()
    with ZipFile(path, "r") as package:
        worksheet_xml = package.read("xl/worksheets/sheet1.xml")
    assert b"<f>1+1</f><v></v>" in worksheet_xml
    worksheet_xml = worksheet_xml.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
    _replace_zip_part(path, "xl/worksheets/sheet1.xml", worksheet_xml)

    result = XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert [block.text for block in result.blocks] == ["Item | Value", "Cached | 2"]
    assert ExtractionWarningCode.FORMULA_WITHOUT_CACHED_VALUE not in result.statistics.warning_codes
    assert "=1+1" not in str(result.to_jsonable())


def test_xlsx_enforces_columns_cells_and_worksheet_limits(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    workbook.active.append(["A", "B", "C"])
    workbook.save(path)
    workbook.close()

    for overrides in (
        {"maximum_spreadsheet_columns": 2, "maximum_csv_columns": 2},
        {"maximum_spreadsheet_cells": 2},
    ):
        with pytest.raises(ExtractionError) as error:
            XlsxExtractor.from_settings(_settings(root)).extract(
                path,
                extraction_limits(**overrides),
            )
        assert error.value.code is ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED


def test_xlsx_enforces_worksheet_and_merged_range_limits(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    workbook.active.append(["A", "B"])
    workbook.active.merge_cells("A1:B1")
    workbook.create_sheet("Second").append(["C"])
    workbook.save(path)
    workbook.close()

    for overrides in (
        {"maximum_worksheets": 1},
        {"maximum_merged_ranges": 0},
    ):
        with pytest.raises(ExtractionError) as error:
            XlsxExtractor.from_settings(_settings(root)).extract(
                path,
                extraction_limits(**overrides),
            )
        assert error.value.code is ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED


def test_xlsx_skips_empty_visible_sheet_and_fails_if_all_are_empty(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    workbook.save(path)
    workbook.close()

    with pytest.raises(ExtractionError) as error:
        XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_NO_EXTRACTABLE_CONTENT


@pytest.mark.parametrize("declared_ref", ["A1:XFD1048576", "A1"])
def test_xlsx_corrects_only_bounded_declared_dimensions(
    tmp_path: Path,
    declared_ref: str,
) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    workbook.active.append(["Header"])
    workbook.active.append(["Value"])
    workbook.save(path)
    workbook.close()
    with ZipFile(path, "r") as package:
        worksheet_xml = package.read("xl/worksheets/sheet1.xml")
    worksheet_xml = worksheet_xml.replace(
        b'ref="A1:A2"',
        f'ref="{declared_ref}"'.encode(),
    )
    _replace_zip_part(path, "xl/worksheets/sheet1.xml", worksheet_xml)

    result = XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert [block.text for block in result.blocks] == ["Header", "Value"]
    assert ExtractionWarningCode.WORKSHEET_DIMENSIONS_CORRECTED in result.statistics.warning_codes


def test_xlsx_rejects_macro_part(tmp_path: Path) -> None:
    root, path = _trusted_path(tmp_path, ".xlsx")
    workbook = Workbook()
    workbook.active.append(["Header"])
    workbook.save(path)
    workbook.close()
    with ZipFile(path, "a", ZIP_DEFLATED) as package:
        package.writestr("xl/vbaProject.bin", b"synthetic")

    with pytest.raises(ExtractionError) as error:
        XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.UNSUPPORTED_DOCUMENT_FORMAT


def test_xlsx_closes_workbook_after_success(tmp_path: Path, monkeypatch) -> None:
    from app.services.ingestion.extractors import xlsx as xlsx_module

    root, path = _trusted_path(tmp_path, ".xlsx")
    source_workbook = Workbook()
    source_workbook.active.append(["Header"])
    source_workbook.save(path)
    source_workbook.close()
    real_workbook = xlsx_module.load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    class WorkbookProxy:
        def __init__(self) -> None:
            self.worksheets = real_workbook.worksheets
            self.closed = False

        def close(self) -> None:
            self.closed = True
            real_workbook.close()

    proxy = WorkbookProxy()
    monkeypatch.setattr(xlsx_module, "load_workbook", lambda **_kwargs: proxy)

    XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert proxy.closed is True


def test_xlsx_close_failure_is_sanitized(tmp_path: Path, monkeypatch) -> None:
    from app.services.ingestion.extractors import xlsx as xlsx_module

    root, path = _trusted_path(tmp_path, ".xlsx")
    source_workbook = Workbook()
    source_workbook.active.append(["Header"])
    source_workbook.save(path)
    source_workbook.close()
    real_workbook = xlsx_module.load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )

    class WorkbookProxy:
        worksheets = real_workbook.worksheets

        def close(self) -> None:
            real_workbook.close()
            raise RuntimeError("private-workbook-path")

    monkeypatch.setattr(xlsx_module, "load_workbook", lambda **_kwargs: WorkbookProxy())

    with pytest.raises(ExtractionError) as error:
        XlsxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_EXTRACTION_FAILED
    assert "private-workbook-path" not in str(error.value)


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_csv_supports_only_allowlisted_delimiters_and_quoted_content(
    tmp_path: Path,
    delimiter: str,
) -> None:
    root, path = _trusted_path(tmp_path, ".csv")
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output, delimiter=delimiter)
        writer.writerow(["Item", "Notes"])
        writer.writerow(["Design", f"quoted{delimiter}value\nnext line"])

    result = CsvExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert result.blocks[0].text == "Item | Notes"
    assert result.blocks[1].text == f"Design | quoted{delimiter}value\nnext line"
    assert result.blocks[1].row_start == 2
    assert result.statistics.extractor_name == CSV_EXTRACTOR_NAME


def test_csv_normalizes_blank_and_duplicate_headers_and_preserves_formula_prefix(
    tmp_path: Path,
) -> None:
    root, path = _trusted_path(tmp_path, ".csv")
    path.write_text("\ufeffName,,name\n=SUM(A1:A2),value,other\n", encoding="utf-8")

    result = CsvExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert result.blocks[0].text == "Name | column_2 | name_2"
    assert result.blocks[1].text.startswith("=SUM(A1:A2)")
    assert set(result.statistics.warning_codes) == {
        ExtractionWarningCode.CSV_HEADER_NORMALIZED,
        ExtractionWarningCode.CSV_DUPLICATE_HEADER_NORMALIZED,
    }


@pytest.mark.parametrize(
    ("content", "overrides", "expected_code"),
    [
        (b"", {}, ExtractionErrorCode.DOCUMENT_NO_EXTRACTABLE_CONTENT),
        (b"\xff\xfeunsafe", {}, ExtractionErrorCode.DOCUMENT_MALFORMED),
        (
            b"A,B,C\n1,2,3\n",
            {"maximum_csv_columns": 2},
            ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED,
        ),
        (
            b"Header\n12345\n",
            {"maximum_csv_field_characters": 4},
            ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED,
        ),
        (
            b"Header\none\ntwo\n",
            {"maximum_rows_per_sheet": 2},
            ExtractionErrorCode.DOCUMENT_EXTRACTION_LIMIT_EXCEEDED,
        ),
    ],
)
def test_csv_rejects_empty_invalid_encoding_and_limits(
    tmp_path: Path,
    content: bytes,
    overrides: dict[str, int],
    expected_code: ExtractionErrorCode,
) -> None:
    root, path = _trusted_path(tmp_path, ".csv")
    path.write_bytes(content)

    with pytest.raises(ExtractionError) as error:
        CsvExtractor.from_settings(_settings(root)).extract(
            path,
            extraction_limits(**overrides),
        )

    assert error.value.code is expected_code


def test_extractors_reject_paths_outside_the_trusted_workspace(tmp_path: Path) -> None:
    root = tmp_path / "processing"
    root.mkdir()
    untrusted = tmp_path / "input.csv"
    untrusted.write_text("A\nB\n", encoding="utf-8")

    with pytest.raises(ExtractionError) as error:
        CsvExtractor.from_settings(_settings(root)).extract(untrusted, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_EXTRACTION_FAILED
    assert str(untrusted) not in str(error.value)


def test_legacy_formats_are_not_advertised() -> None:
    from app.api.routes.documents import PROJECT_EXTENSIONS
    from app.api.routes.templates import TEMPLATE_EXTENSIONS

    assert ".doc" not in PROJECT_EXTENSIONS
    assert ".xls" not in PROJECT_EXTENSIONS
    assert ".xls" not in TEMPLATE_EXTENSIONS
    create_wbs_source = (
        Path(__file__).resolve().parents[2] / "WBS_UI" / "src" / "views" / "CreateWbsScreen.tsx"
    ).read_text(encoding="utf-8")
    assert "['.xlsx', '.xls'" not in create_wbs_source


def test_unexpected_parser_error_is_sanitized(tmp_path: Path, monkeypatch, caplog) -> None:
    from app.services.ingestion.extractors import docx as docx_module

    root, path = _trusted_path(tmp_path, ".docx")
    document = Document()
    document.add_paragraph("Synthetic scope")
    document.save(path)
    sensitive_marker = "private-path filename storage-key endpoint document-content"

    def fail_parser(*_args: object, **_kwargs: object) -> None:
        raise RuntimeError(sensitive_marker)

    monkeypatch.setattr(docx_module, "Document", fail_parser)
    with pytest.raises(ExtractionError) as error:
        DocxExtractor.from_settings(_settings(root)).extract(path, extraction_limits())

    assert error.value.code is ExtractionErrorCode.DOCUMENT_EXTRACTION_FAILED
    assert sensitive_marker not in str(error.value)
    assert sensitive_marker not in caplog.text
    assert str(path) not in caplog.text


def _replace_zip_part(path: Path, target_name: str, replacement: bytes) -> None:
    temporary = path.with_suffix(".replacement")
    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as output:
        for entry in source.infolist():
            data = replacement if entry.filename == target_name else source.read(entry.filename)
            output.writestr(entry, data)
    temporary.replace(path)
